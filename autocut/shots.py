# -*- coding: utf-8 -*-
"""分镜执行表的解析。

单独成模块的原因：这一层不依赖 pymiere，也不依赖 Premiere 在不在跑，
因此可以在任何机器（含 CI）上测试。PR 相关的逻辑都在 scripts/ 里。
"""
import re


def parse_tc(tc):
    """"0:12-0:19" -> 12.0 ；"1:04-1:10" -> 64.0

    返回**时间线起点（秒）**。分镜表的「时间码」列记的是成片时间轴上的区间，
    我们只需要起点来决定片段落在时间线的什么位置。
    解析失败返回 None（调用方据此跳过该镜头）。
    """
    m = re.match(r"\s*(\d+):(\d+(?:\.\d+)?)\s*-\s*(\d+):(\d+(?:\.\d+)?)", tc or "")
    if not m:
        return None
    return int(m.group(1)) * 60 + float(m.group(2))


def first_clip(cell):
    """从素材单元格取出第一个素材名。

    "CLIP_6008 / 6030 / 6068"      -> "CLIP_6008"
    "DJI_0085 / DJI_0086"          -> "DJI_0085"
    返回 None 表示这一镜没有素材（字幕卡 / 黑场），时间线上留空。
    """
    if not cell:
        return None
    tok = cell.strip().split("/")[0].strip()
    tok = tok.replace("\\", " ").strip()
    return tok or None


def to_float(v):
    """把单元格转成 float；转不了返回 None。"""
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_row(row, columns, index=None):
    """把一行表数据解析成一个镜头 dict。

    columns 来自 config.json 的 columns 字段（0-based 列号）。
    """
    _r = list(row) + [None] * 24

    def take(key):
        i = columns.get(key)
        if i is None or i >= len(_r):
            return None
        return _r[i]

    clip = first_clip(take("media"))
    shot = {
        "no": take("no"),
        "tl": parse_tc(take("tc")),
        "clip": clip,
        "in": to_float(take("src_in")),
        "out": to_float(take("src_out")),
        "desc": take("desc"),
        "snd": take("snd"),
        "note": take("note") or "",
    }
    return shot


def load_shots(path, sequences, columns):
    """读取分镜表，返回 [(中文序列名, ASCII 序列名, [镜头...]), ...]

    需要 openpyxl。这里延迟导入，让不读表的调用方（例如纯配置校验）无需安装。
    """
    from openpyxl import load_workbook

    header_rows = columns.get("header_rows", 2)
    wb = load_workbook(path, data_only=True)
    out = []
    for spec in sequences:
        sheet = spec["sheet"]
        if sheet not in wb.sheetnames:
            raise KeyError(
                "分镜表里没有工作表 %r；实际有：%s" % (sheet, ", ".join(wb.sheetnames)))
        ws = wb[sheet]
        shots = []
        for row in ws.iter_rows(min_row=header_rows + 1, values_only=True):
            if not row or row[0] is None:
                continue
            shots.append(parse_row(row, columns))
        out.append((spec["title"], spec["ascii_name"], shots))
    return out
